from datetime import datetime
from io import BytesIO
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from astraea.reference.team_names import TEAM_NAMES
from astraea.storage import load_json, update_json

FILES = {'match': 'scouting_data.json', 'pit': 'pit_data.json'}

COLLECT_SOURCES = {'neutral': 'Neutral Zone', 'outpost': 'Outpost', 'depot': 'Depot', 'preloaded': 'Preloaded'}
AUTO_STATUSES = {'active': 'Active', 'no_move': "Didn't move", 'moved_no_score': 'Moved, no score'}
ROLES = {'offense': 'Offense', 'defense': 'Defense', 'feeder': 'Feeder', 'mix': 'Mix'}
CLIMBS = {'none': ('None', 0), 'level1': ('L1', 10), 'level2': ('L2', 20), 'level3': ('L3', 30)}
DRIVEBASES = {'tank': 'Tank', 'mecanum': 'Mecanum', 'swerve': 'Swerve', 'west_coast': 'West Coast', 'other': 'Other'}
AUTO_CLIMB_POINTS = 15

COLUMNS = {
    'match': [
        ('team', 'Team'),
        ('team_name', 'Team Name'),
        ('event', 'Event'),
        ('match', 'Match'),
        ('scouter', 'Scouter'),
        ('auto_status', 'Auto Status'),
        ('auto_fuel_scored', 'Auto Fuel Scored'),
        ('auto_fuel_missed', 'Auto Fuel Missed'),
        ('auto_collect', 'Auto Collected From'),
        ('auto_climb', 'Auto Climb'),
        ('teleop_moved', 'Teleop Moved'),
        ('teleop_fuel_scored', 'Teleop Fuel Scored'),
        ('teleop_fuel_missed', 'Teleop Fuel Missed'),
        ('can_cross_bump', 'Crosses Bump'),
        ('can_cross_trench', 'Crosses Trench'),
        ('teleop_collect', 'Teleop Collected From'),
        ('robot_role', 'Role'),
        ('offense_rating', 'Offense Rating'),
        ('defense_rating', 'Defense Rating'),
        ('climb', 'Tower Climb'),
        ('climb_result', 'Climb Result'),
        ('partial_match', 'Partial Match'),
        ('points', 'Est. Points'),
        ('notes', 'Notes'),
        ('submitted_at', 'Entered'),
    ],
    'pit': [
        ('team', 'Team'),
        ('team_name', 'Team Name'),
        ('event', 'Event'),
        ('scouter', 'Scouter'),
        ('drivebase', 'Drivebase'),
        ('avg_cycle_time', 'Avg Cycle Time (s)'),
        ('notes', 'Notes'),
        ('submitted_at', 'Entered'),
    ],
}

SHEET_TITLES = {'match': 'Match Scouting', 'pit': 'Pit Scouting'}


class ValidationError(ValueError):
    pass


def _number(value, maximum=None, required_field=None):
    text = str(value if value is not None else '').strip()
    if not text:
        if required_field:
            raise ValidationError(f'{required_field} is required')
        return None
    try:
        number = int(float(text))
    except ValueError:
        raise ValidationError(f'"{text}" is not a number')
    if number < 0 or (maximum is not None and number > maximum):
        raise ValidationError(f'{number} is out of range')
    return number


def _sources(values):
    return [v for v in values or [] if v in COLLECT_SOURCES]


def _text(value):
    return str(value or '').strip()


def team_name(team):
    return TEAM_NAMES.get(str(team), '')


def clean_match(data):
    auto_status = data.get('auto_status') if data.get('auto_status') in AUTO_STATUSES else 'active'
    auto_active = auto_status == 'active'
    teleop_moved = bool(data.get('teleop_moved', True))
    role = data.get('robot_role') if data.get('robot_role') in ROLES else ''
    climb = data.get('climb') if data.get('climb') in CLIMBS else 'none'

    return {
        'event': _text(data.get('event')),
        'scouter': _text(data.get('scouter')),
        'team': _number(data.get('team'), required_field='Team number'),
        'match': _number(data.get('match'), required_field='Match number'),
        'auto_status': auto_status,
        'auto_fuel_scored': (_number(data.get('auto_fuel_scored')) or 0) if auto_active else 0,
        'auto_fuel_missed': (_number(data.get('auto_fuel_missed')) or 0) if auto_active else 0,
        'auto_collect': _sources(data.get('auto_collect')) if auto_active else [],
        'auto_climb': auto_active and bool(data.get('auto_climb')),
        'teleop_moved': teleop_moved,
        'teleop_fuel_scored': (_number(data.get('teleop_fuel_scored')) or 0) if teleop_moved else 0,
        'teleop_fuel_missed': (_number(data.get('teleop_fuel_missed')) or 0) if teleop_moved else 0,
        'can_cross_bump': bool(data.get('can_cross_bump')),
        'can_cross_trench': bool(data.get('can_cross_trench')),
        'teleop_collect': _sources(data.get('teleop_collect')) if teleop_moved else [],
        'robot_role': role,
        'offense_rating': _number(data.get('offense_rating'), 5) if role in ('offense', 'mix') else None,
        'defense_rating': _number(data.get('defense_rating'), 5) if role in ('defense', 'mix') else None,
        'climb': climb,
        'climb_success': climb != 'none' and bool(data.get('climb_success')),
        'partial_match': bool(data.get('partial_match')),
        'notes': _text(data.get('notes')),
    }


def clean_pit(data):
    try:
        cycle = round(float(data.get('avg_cycle_time')), 1)
    except (TypeError, ValueError):
        cycle = None
    drivebase = data.get('drivebase') if data.get('drivebase') in DRIVEBASES else ''
    return {
        'event': _text(data.get('event')),
        'scouter': _text(data.get('scouter')),
        'team': _number(data.get('team'), required_field='Team number'),
        'drivebase': drivebase,
        'avg_cycle_time': cycle,
        'notes': _text(data.get('notes')),
    }


CLEANERS = {'match': clean_match, 'pit': clean_pit}


def points(entry):
    climb_points = CLIMBS[entry['climb']][1] if entry['climb_success'] else 0
    return (entry['auto_fuel_scored'] + entry['teleop_fuel_scored']
            + (AUTO_CLIMB_POINTS if entry['auto_climb'] else 0) + climb_points)


def list_entries(kind):
    return load_json(FILES[kind], [])


def get_entry(kind, entry_id):
    return next((e for e in list_entries(kind) if e['id'] == entry_id), None)


def _same_slot(kind, a, b):
    keys = ('event', 'team', 'match') if kind == 'match' else ('event', 'team')
    return all(a.get(k) == b.get(k) for k in keys)


def find_duplicate(kind, entry, exclude_id=None):
    return next((e for e in list_entries(kind) if e['id'] != exclude_id and _same_slot(kind, e, entry)), None)


def save_entry(kind, data, entry_id=None):
    entry = CLEANERS[kind](data)
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with update_json(FILES[kind], []) as entries:
        if entry_id is None:
            entry = {'id': uuid4().hex[:12], 'submitted_at': now, **entry}
            entries.append(entry)
            return entry
        for i, existing in enumerate(entries):
            if existing['id'] == entry_id:
                entries[i] = entry = {'id': entry_id, 'submitted_at': existing['submitted_at'], 'updated_at': now, **entry}
                return entry
    return None


def delete_entry(kind, entry_id):
    with update_json(FILES[kind], []) as entries:
        before = len(entries)
        entries[:] = [e for e in entries if e['id'] != entry_id]
        return len(entries) < before


def _labels(values):
    return ', '.join(COLLECT_SOURCES[v] for v in values)


def display_row(kind, e):
    row = {**e, 'team_name': team_name(e['team'])}
    if kind == 'match':
        level = CLIMBS[e['climb']][0]
        row.update({
            'auto_status': AUTO_STATUSES[e['auto_status']],
            'auto_collect': _labels(e['auto_collect']),
            'teleop_collect': _labels(e['teleop_collect']),
            'robot_role': ROLES.get(e['robot_role'], ''),
            'climb': level,
            'climb_result': '' if level == 'None' else ('Success' if e['climb_success'] else 'Failed'),
            'points': points(e),
        })
    else:
        row['drivebase'] = DRIVEBASES.get(e['drivebase'], '')
    return {key: ('Yes' if v is True else 'No' if v is False else '' if v is None else v)
            for key, v in row.items()}


def table(kind):
    entries = sorted(list_entries(kind), key=lambda e: (e['team'], e.get('match', 0), e['submitted_at']))
    rows = [display_row(kind, e) for e in entries]
    return {'columns': [{'key': k, 'label': label} for k, label in COLUMNS[kind]], 'rows': rows}


HEADER_FILL = PatternFill('solid', fgColor='0B1F4D')
HEADER_FONT = Font(bold=True, color='FFFFFF')
PARTIAL_FILL = PatternFill('solid', fgColor='FFF3B0')


def _write_sheet(ws, kind):
    columns = COLUMNS[kind]
    rows = table(kind)['rows']
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in rows:
        ws.append([row.get(key, '') for key, _ in columns])
        if row.get('partial_match') == 'Yes':
            for cell in ws[ws.max_row]:
                cell.fill = PARTIAL_FILL

    for i, (key, label) in enumerate(columns, 1):
        letter = get_column_letter(i)
        longest = max([len(label)] + [len(str(r.get(key, ''))) for r in rows])
        ws.column_dimensions[letter].width = min(longest + 3, 60)
        if key == 'notes':
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical='top')

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def export_workbook():
    wb = Workbook()
    for i, kind in enumerate(FILES):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = SHEET_TITLES[kind]
        _write_sheet(ws, kind)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
